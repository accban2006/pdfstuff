from queue import Queue
from threading import Event

import pytest
from docx import Document

from pdftomd import jobs
from pdftomd.errors import ConversionError


def test_snapshot_collisions_and_mixed_batch(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    good = source / "Tài liệu.md"
    good.write_text("# Tiếng Việt **đậm**", encoding="utf-8")
    (source / "bad.md").write_bytes(b"\xff")
    (source / "~$lock.md").write_text("ignored")
    (source / "~temp.md").write_text("ignored")
    (source / "child").mkdir()
    (source / "child/nested.md").write_text("ignored")
    plan = jobs.prepare(str(source), str(source), "md_word")
    (source / "later.md").write_text("not in snapshot")
    original = good.read_bytes()
    existing = source / "TÀI LIỆU.DOCX"
    existing.write_bytes(b"keep me")
    events = Queue()
    jobs.run(plan, Event(), events)
    results = [event.result for event in list(events.queue) if event.kind == "result"]
    assert [r.status for r in results] == ["error", "success"]
    assert results[1].output.name == "Tài liệu (1).docx"
    assert existing.read_bytes() == b"keep me"
    assert good.read_bytes() == original
    assert not list(source.glob(".pdftomd-*"))
    assert list(events.queue)[-1].kind == "done"
    assert [e.current for e in events.queue if e.kind == "result"] == [1, 2]


@pytest.mark.parametrize(
    "source,mode,key",
    [
        ("absent.md", "md_word", "source_missing"),
        ("legacy.doc", "word_md", "legacy_doc"),
        ("wrong.txt", "md_word", "wrong_extension"),
    ],
)
def test_preflight(tmp_path, source, mode, key):
    path = tmp_path / source
    if source != "absent.md":
        path.write_text("text")
    with pytest.raises(ConversionError) as error:
        jobs.prepare(str(path), str(tmp_path / "out"), mode)
    assert error.value.notice.key == key


def test_empty_folder_missing_dependencies(tmp_path, monkeypatch):
    with pytest.raises(ConversionError, match="no_files"):
        jobs.prepare(str(tmp_path), str(tmp_path / "out"), "md_word")
    monkeypatch.setattr(jobs.word_to_pdf, "find_libreoffice", lambda: None)
    with pytest.raises(ConversionError, match="missing_lo"):
        jobs.check_dependencies("word_pdf")
    monkeypatch.setattr(
        jobs, "require_arial", lambda: (_ for _ in ()).throw(ConversionError("missing_arial"))
    )
    with pytest.raises(ConversionError, match="missing_arial"):
        jobs.check_dependencies("md_word")


def test_cancel_during_file_never_publishes(tmp_path, monkeypatch):
    for name in ("a.md", "b.md"):
        (tmp_path / name).write_text("valid")
    plan = jobs.prepare(str(tmp_path), str(tmp_path / "out"), "md_word")
    cancel = Event()

    def converting(source, destination, event):
        destination.write_bytes(b"incomplete")
        event.set()
        return []

    monkeypatch.setattr(jobs.markdown_to_docx, "convert", converting)
    events = Queue()
    jobs.run(plan, cancel, events)
    assert [e.result.status for e in events.queue if e.kind == "result"] == [
        "cancelled",
        "cancelled",
    ]
    assert list(plan.destination.iterdir()) == []


def test_output_permission_and_publish_cleanup(tmp_path, monkeypatch):
    source = tmp_path / "a.md"
    source.write_text("valid")
    monkeypatch.setattr(
        jobs.tempfile,
        "TemporaryFile",
        lambda **kw: (_ for _ in ()).throw(PermissionError("denied")),
    )
    with pytest.raises(PermissionError):
        jobs.prepare(str(source), str(tmp_path / "out"), "md_word")
    staged = tmp_path / "staged.docx"
    staged.write_bytes(b"complete")
    monkeypatch.setattr(jobs.os, "replace", lambda *args: (_ for _ in ()).throw(OSError()))
    with pytest.raises(OSError):
        jobs.publish(staged, tmp_path / "out", "final", ".docx")
    assert not (tmp_path / "out/final.docx").exists()


def test_long_unicode_path(tmp_path):
    folder = tmp_path / ("thư mục " * 4).rstrip() / ("English folder " * 2).rstrip()
    folder.mkdir(parents=True)
    source = folder / "Tiếng Việt có dấu và khoảng trắng.md"
    source.write_text("English / Tiếng Việt", encoding="utf-8")
    plan = jobs.prepare(str(source), str(folder / "out"), "md_word")
    events = Queue()
    jobs.run(plan, Event(), events)
    result = next(e.result for e in events.queue if e.kind == "result")
    assert result.status == "success"
    assert result.output.is_file()


def test_cancel_keeps_completed_output(tmp_path):
    for name in ("a.md", "b.md"):
        (tmp_path / name).write_text("valid")
    plan = jobs.prepare(str(tmp_path), str(tmp_path / "out"), "md_word")
    cancel = Event()

    class CancellingQueue(Queue):
        def put(self, event, *args, **kwargs):
            super().put(event, *args, **kwargs)
            if event.kind == "result" and event.current == 1:
                cancel.set()

    events = CancellingQueue()
    jobs.run(plan, cancel, events)
    results = [e.result for e in events.queue if e.kind == "result"]
    assert [r.status for r in results] == ["success", "cancelled"]
    assert results[0].output.is_file()
    assert not (plan.destination / "b.docx").exists()


@pytest.mark.parametrize("words", ["English document", "Tài liệu tiếng Việt Đặng"])
@pytest.mark.parametrize("outcome", ["success", "error", "cancel"])
def test_markdown_pdf_uses_shared_docx_and_publishes_safely(tmp_path, monkeypatch, words, outcome):
    source = tmp_path / "Tài liệu.md"
    source.write_text(f"# {words}\n\nNormal **bold**", encoding="utf-8-sig")
    original = source.read_bytes()
    monkeypatch.setattr(jobs.word_to_pdf, "find_libreoffice", lambda: source)
    plan = jobs.prepare(str(source), str(tmp_path / "out"), "md_pdf")
    existing = plan.destination / "TÀI LIỆU.PDF"
    existing.write_bytes(b"existing file")
    intermediate_paths = []

    def export(intermediate, staged, cancel, *, arial):
        intermediate_paths.append(intermediate)
        assert arial is True
        doc = Document(intermediate)
        assert doc.paragraphs[0].text == words
        assert doc.paragraphs[0].style.name == "Heading 1"
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                assert run.font.name == "Arial"
                assert run.font.size.pt == (15 if run.bold else 11)
                assert str(run.font.color.rgb) == "000000"
        if outcome == "error":
            staged.write_bytes(b"incomplete")
            raise ConversionError("lo_failed", "1")
        staged.write_bytes(b"mock PDF output")
        if outcome == "cancel":
            cancel.set()

    monkeypatch.setattr(jobs.word_to_pdf, "convert", export)
    events = Queue()
    jobs.run(plan, Event(), events)
    result = next(e.result for e in events.queue if e.kind == "result")
    assert result.status == {"success": "warning", "error": "error", "cancel": "cancelled"}[outcome]
    if outcome == "success":
        assert result.output.name == "Tài liệu (1).pdf"
        assert [notice.key for notice in result.notices] == ["pdf_unverified"]
    else:
        assert result.output is None
        assert list(plan.destination.iterdir()) == [existing]
    assert len(intermediate_paths) == 1
    assert not intermediate_paths[0].exists()
    assert not list(plan.destination.glob(".pdftomd-*"))
    assert existing.read_bytes() == b"existing file"
    assert source.read_bytes() == original
    assert list(events.queue)[-1].kind == "done"
