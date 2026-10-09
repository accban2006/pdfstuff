from pathlib import Path
from threading import Event

import pytest
from docx import Document
from PIL import Image

from pdftomd.converters.markdown_to_docx import convert
from pdftomd.errors import Cancelled, ConversionError


@pytest.mark.parametrize("words", ["English document © ± €", "Tiếng Việt: chuyển đổi Đặng © ± €"])
def test_typography_and_structure(tmp_path: Path, words: str):
    source = tmp_path / "Tài liệu mẫu.md"
    source.write_text(
        "\n".join(f"{'#' * level} {words}" for level in range(1, 7))
        + f"\n\nNormal {words} **bold** *italic* [link](https://example.com)\n\n"
        + "3. **list bold**\n4. item\n   - nested\n\n"
        + "| Header | Other |\n| --- | ---: |\n| **table bold** | value |\n\n"
        + "```python\n  a = 1\n\tb = 2\n```\n",
        encoding="utf-8-sig",
    )
    target = tmp_path / "result.docx"
    assert convert(source, target, Event()) == []
    doc = Document(target)
    for index, paragraph in enumerate(doc.paragraphs[:6], 1):
        assert paragraph.style.name == f"Heading {index}"
        assert paragraph.text == words
        assert paragraph.runs[0].font.size.pt == 15
    all_paragraphs = doc.paragraphs + [
        p for row in doc.tables[0].rows for cell in row.cells for p in cell.paragraphs
    ]
    for paragraph in all_paragraphs:
        for run in paragraph.runs:
            assert run.font.name == "Arial"
            assert str(run.font.color.rgb) == "000000"
            assert run.font.size.pt == (15 if run.bold else 11)
    paragraph = next(p for p in doc.paragraphs if p.text == "list bold")
    assert paragraph._p.pPr.numPr.numId.val > 0
    assert doc.tables[0].cell(1, 0).text == "table bold"
    assert any(p.text == "  a = 1\n\tb = 2\n" for p in doc.paragraphs)
    assert any(rel.target_ref == "https://example.com" for rel in doc.part.rels.values())
    assert doc.sections[0].page_width.cm == pytest.approx(21, abs=0.01)


def test_images_and_html(tmp_path):
    Image.new("RGB", (200, 100), "blue").save(tmp_path / "ảnh.png")
    source = tmp_path / "test.md"
    source.write_text("![ảnh](ảnh.png)\n\n<script>alert(1)</script>", encoding="utf-8")
    target = tmp_path / "test.docx"
    notices = convert(source, target, Event())
    assert [notice.key for notice in notices] == ["html_literal"]
    doc = Document(target)
    assert len(doc.inline_shapes) == 1
    assert doc.inline_shapes[0].width / doc.inline_shapes[0].height == pytest.approx(2)
    assert "<script>alert(1)</script>" in "".join(p.text for p in doc.paragraphs)


@pytest.mark.parametrize(
    "text,key",
    [
        ("![x](missing.png)", "missing_image"),
        ("![x](https://example.com/a.png)", "remote_image"),
        ("Unsupported 😀", "missing_glyph"),
    ],
)
def test_invalid_resources(tmp_path, text, key):
    source = tmp_path / "test.md"
    source.write_text(text, encoding="utf-8")
    with pytest.raises(ConversionError) as error:
        convert(source, tmp_path / "test.docx", Event())
    assert error.value.notice.key == key
    assert not (tmp_path / "test.docx").exists()


def test_encoding_and_cancel(tmp_path):
    source = tmp_path / "bad.md"
    source.write_bytes(b"\xff")
    with pytest.raises(ConversionError, match="invalid_utf8"):
        convert(source, tmp_path / "out.docx", Event())
    source.write_text("hello", encoding="utf-8")
    cancel = Event()
    cancel.set()
    with pytest.raises(Cancelled):
        convert(source, tmp_path / "out.docx", cancel)


def test_inline_code_preserves_whitespace(tmp_path):
    source = tmp_path / "code.md"
    source.write_text("`  space  ` and ``a\nb``", encoding="utf-8")
    target = tmp_path / "code.docx"
    convert(source, target, Event())
    assert Document(target).paragraphs[0].text == "  space   and a\nb"


@pytest.mark.parametrize("first", ["# heading", "> quoted", "```\n   code\n```"])
def test_list_items_with_nonparagraph_first_block(tmp_path, first):
    source = tmp_path / "list.md"
    # Indent continuation lines to keep the fenced block inside the list item.
    first = first.replace("\n", "\n   ")
    source.write_text(f"3. {first}\n\n4. second", encoding="utf-8")
    target = tmp_path / "list.docx"
    convert(source, target, Event())
    doc = Document(target)
    paragraphs = [p for p in doc.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]
    assert len(paragraphs) == 2
    assert paragraphs[0]._p.pPr.numPr.numId.val == paragraphs[1]._p.pPr.numPr.numId.val


def test_footnote_definition_never_drops_text(tmp_path):
    source = tmp_path / "footnote.md"
    source.write_text("Body[^a]\n\n[^a]: Important footnote content", encoding="utf-8")
    target = tmp_path / "footnote.docx"
    with pytest.raises(ConversionError, match="footnote"):
        convert(source, target, Event())
    assert not target.exists()
