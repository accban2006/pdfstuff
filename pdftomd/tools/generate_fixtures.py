"""Regenerate synthetic fixtures using developer-only reportlab and python-docx."""

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1] / "tests" / "fixtures"
ROOT.mkdir(parents=True, exist_ok=True)
pdfmetrics.registerFont(TTFont("ArialFixture", "C:/Windows/Fonts/arial.ttf"))

for code, words in (
    ("en", "English document: café © ± € 123"),
    ("vi", "Tiếng Việt: chuyển đổi Đặng © ± € 123"),
):
    doc = Document()
    doc.add_heading(words, 1)
    doc.add_paragraph("Normal " + words)
    doc.add_paragraph("First item", style="List Bullet")
    doc.add_paragraph("Second item", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Name", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = words, "42"
    doc.save(ROOT / f"{code}.docx")
    pdf = canvas.Canvas(str(ROOT / f"{code}.pdf"), pagesize=(595.28, 841.89))
    pdf.setFont("ArialFixture", 15)
    pdf.drawString(55, 790, words)
    pdf.setFont("ArialFixture", 11)
    pdf.drawString(55, 755, "Normal " + words)
    pdf.drawString(55, 725, "1. First item")
    pdf.drawString(55, 705, "2. Second item")
    for x in (55, 300, 530):
        pdf.line(x, 610, x, 675)
    for y in (610, 642, 675):
        pdf.line(55, y, 530, y)
    for x, y, value in (
        (65, 655, "Name"),
        (310, 655, "Value"),
        (65, 622, "Tiếng Việt" if code == "vi" else "English"),
        (310, 622, "42"),
    ):
        pdf.drawString(x, y, value)
    pdf.save()

image = Image.new("RGB", (600, 240), "white")
draw = ImageDraw.Draw(image)
draw.text((30, 50), "Scanned text / Van ban scan", fill="black")
image.save(ROOT / "local.png")
pdf = canvas.Canvas(str(ROOT / "scan.pdf"))
pdf.drawImage(str(ROOT / "local.png"), 50, 600, width=450, height=180)
pdf.save()

doc = Document()
doc.sections[0].header.paragraphs[0].text = "HEADER / ĐẦU TRANG"
doc.sections[0].footer.paragraphs[0].text = "FOOTER / CUỐI TRANG"
paragraph = doc.add_paragraph()
run = paragraph.add_run("Original Times New Roman 18 / Tiếng Việt")
run.font.name, run.font.size = "Times New Roman", Pt(18)
doc.add_picture(str(ROOT / "local.png"), width=Inches(3))
table = doc.add_table(rows=2, cols=2)
table.style = "Table Grid"
table.cell(0, 0).text, table.cell(0, 1).text = "Name", "Value"
table.cell(1, 0).text, table.cell(1, 1).text = "Tiếng Việt", "42"
doc.add_page_break()
doc.add_paragraph("PAGE TWO / TRANG HAI")
doc.save(ROOT / "word-layout.docx")
