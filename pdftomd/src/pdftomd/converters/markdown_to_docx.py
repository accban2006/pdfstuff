"""Render CommonMark tokens and tables into one shared DOCX representation."""

from pathlib import Path
from threading import Event
from urllib.parse import unquote, urlsplit

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu
from docx.text.paragraph import Paragraph
from markdown_it import MarkdownIt
from markdown_it.rules_block.state_block import StateBlock
from markdown_it.rules_inline.backticks import backtick
from markdown_it.rules_inline.state_inline import StateInline
from markdown_it.token import Token
from markdown_it.tree import SyntaxTreeNode
from PIL import Image

from pdftomd.errors import Cancelled, ConversionError, Notice
from pdftomd.styles import (
    CONTENT_WIDTH_CM,
    add_numbering,
    check_glyphs,
    format_font,
    new_document,
)


def preserving_backtick(state: StateInline, silent: bool) -> bool:
    start, count = state.pos, len(state.tokens)
    matched = backtick(state, silent)
    if matched and not silent and len(state.tokens) > count:
        token = state.tokens[-1]
        if token.type == "code_inline":
            length = len(token.markup)
            token.content = state.src[start + length : state.pos - length]
    return matched


def reject_footnote_definition(state: StateBlock, start: int, end: int, silent: bool) -> bool:
    offset = state.bMarks[start] + state.tShift[start]
    line = state.src[offset : state.eMarks[start]]
    if line.startswith("[^") and "]:" in line:
        raise ConversionError("unsupported", "footnote")
    return False


def parser() -> MarkdownIt:
    # HTML is parsed only to preserve it literally and report a warning; never executed.
    backend = MarkdownIt("commonmark", {"html": True, "typographer": False}).enable(
        ["table", "strikethrough"]
    )
    backend.inline.ruler.at("backticks", preserving_backtick)
    # CommonMark treats [^note]: as a reference definition and would drop the text.
    backend.block.ruler.before("reference", "unsupported_footnote", reject_footnote_definition)
    return backend


class Renderer:
    def __init__(self, source: Path, cancel: Event) -> None:
        self.source = source
        self.cancel = cancel
        self.document = new_document()
        self.notices: list[Notice] = []

    def checkpoint(self) -> None:
        if self.cancel.is_set():
            raise Cancelled()

    def inline(
        self,
        paragraph: Paragraph,
        tokens: list[Token],
        *,
        heading: bool = False,
        width_cm: float = CONTENT_WIDTH_CM,
    ) -> None:
        strong = italic = strike = 0
        hyperlink = None
        for token in tokens:
            self.checkpoint()
            kind = token.type
            if kind in ("strong_open", "strong_close"):
                strong += 1 if token.nesting == 1 else -1
            elif kind in ("em_open", "em_close"):
                italic += 1 if token.nesting == 1 else -1
            elif kind in ("s_open", "s_close"):
                strike += 1 if token.nesting == 1 else -1
            elif kind == "link_open":
                hyperlink = OxmlElement("w:hyperlink")
                hyperlink.set(
                    qn("r:id"),
                    paragraph.part.relate_to(token.attrGet("href"), RT.HYPERLINK, is_external=True),
                )
                paragraph._p.append(hyperlink)
            elif kind == "link_close":
                hyperlink = None
            elif kind == "image":
                self.image(paragraph, token.attrGet("src") or "", token.content, width_cm)
                if hyperlink is not None:
                    self.notices.append(Notice("image_link"))
            elif kind in ("text", "code_inline", "html_inline", "softbreak", "hardbreak"):
                text = "\n" if kind in ("softbreak", "hardbreak") else token.content
                if kind == "text" and ("$$" in text or "[^" in text):
                    self.notices.append(Notice("markdown_extension"))
                check_glyphs(text)
                run = paragraph.add_run(text)
                format_font(run.font, strong=heading or strong > 0, italic=italic > 0)
                run.font.strike = strike > 0
                if hyperlink is not None:
                    hyperlink.append(run._r)
                if kind == "html_inline":
                    self.notices.append(Notice("html_literal"))
            else:
                raise ConversionError("unsupported", kind)

    def image(self, paragraph: Paragraph, reference: str, alt: str, width_cm: float) -> None:
        url = urlsplit(reference)
        if url.scheme or url.netloc:
            raise ConversionError("remote_image", reference)
        path = (self.source.parent / unquote(url.path)).resolve()
        if not path.is_file():
            raise ConversionError("missing_image", reference)
        try:
            with Image.open(path) as picture:
                pixels_w, pixels_h = picture.size
                dpi = picture.info.get("dpi", (96, 96))[0] or 96
                width = min(Cm(width_cm), Emu(pixels_w / dpi * 914400))
                width = min(width, Emu(Cm(24) * pixels_w / pixels_h))
            shape = paragraph.add_run().add_picture(str(path), width=width)
            shape._inline.docPr.set("descr", alt)
        except Exception as exc:
            raise ConversionError("invalid_image", reference) from exc

    def blocks(self, nodes: list[SyntaxTreeNode], *, depth: int = 0, quote: bool = False) -> None:
        for node in nodes:
            self.checkpoint()
            kind = node.type
            if kind in ("paragraph", "heading"):
                heading = kind == "heading"
                style = f"Heading {node.tag[1:]}" if heading else "Quote" if quote else None
                paragraph = self.document.add_paragraph(style=style)
                if depth:
                    paragraph.paragraph_format.left_indent = Cm(0.7 * depth)
                for child in node.children:
                    if child.type == "inline":
                        self.inline(paragraph, child.token.children or [], heading=heading)
            elif kind in ("bullet_list", "ordered_list"):
                self.list_block(node, depth, quote)
            elif kind == "blockquote":
                self.blocks(node.children, depth=depth, quote=True)
            elif kind == "table":
                self.table(node)
            elif kind in ("fence", "code_block", "html_block"):
                check_glyphs(node.content)
                paragraph = self.document.add_paragraph()
                format_font(paragraph.add_run(node.content).font)
                if kind == "html_block":
                    self.notices.append(Notice("html_literal"))
            elif kind == "hr":
                paragraph = self.document.add_paragraph()
                borders = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                for key, value in (("val", "single"), ("sz", "4"), ("color", "000000")):
                    bottom.set(qn(f"w:{key}"), value)
                borders.append(bottom)
                paragraph._p.get_or_add_pPr().append(borders)
            else:
                raise ConversionError("unsupported", kind)

    def list_block(self, node: SyntaxTreeNode, depth: int, quote: bool) -> None:
        if depth > 8:
            raise ConversionError("unsupported", "list nesting > 9")
        ordered = node.type == "ordered_list"
        number = int(node.attrs.get("start", 1))
        number_id = add_numbering(self.document, ordered=ordered, start=number, depth=depth)

        def mark(paragraph: Paragraph) -> None:
            paragraph.paragraph_format.left_indent = Cm(0.7 * (depth + 1))
            paragraph.paragraph_format.first_line_indent = Cm(-0.5)
            numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
            numbering.get_or_add_ilvl().val = depth
            numbering.get_or_add_numId().val = number_id

        for item in node.children:
            first = True
            for block in item.children:
                if block.type == "paragraph":
                    paragraph = self.document.add_paragraph()
                    paragraph.paragraph_format.left_indent = Cm(0.7 * (depth + 1))
                    paragraph.paragraph_format.first_line_indent = Cm(-0.5) if first else Cm(0)
                    if first:
                        mark(paragraph)
                    for inline in block.children:
                        self.inline(
                            paragraph,
                            inline.token.children or [],
                            width_cm=CONTENT_WIDTH_CM - 0.7 * (depth + 1),
                        )
                    first = False
                else:
                    if first and block.type in ("table", "bullet_list", "ordered_list"):
                        mark(self.document.add_paragraph())
                        first = False
                    previous_count = len(self.document.paragraphs)
                    self.blocks([block], depth=depth + 1, quote=quote)
                    if first:
                        mark(self.document.paragraphs[previous_count])
                    first = False
            if first:
                mark(self.document.add_paragraph())
            number += 1

    def table(self, node: SyntaxTreeNode) -> None:
        rows = [row for group in node.children for row in group.children]
        columns = len(rows[0].children)
        table = self.document.add_table(rows=0, cols=columns)
        table.style = "Table Grid"
        table.autofit = False
        for column in table.columns:
            column.width = Cm(CONTENT_WIDTH_CM / columns)
        for row_node in rows:
            cells = table.add_row().cells
            for cell, cell_node in zip(cells, row_node.children, strict=True):
                cell.width = Cm(CONTENT_WIDTH_CM / columns)
                paragraph = cell.paragraphs[0]
                alignment = cell_node.attrs.get("style", "")
                if "right" in alignment:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif "center" in alignment:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for inline in cell_node.children:
                    self.inline(
                        paragraph,
                        inline.token.children or [],
                        width_cm=max(0.5, CONTENT_WIDTH_CM / columns - 0.4),
                    )
            if row_node.children[0].type == "th":
                repeat = OxmlElement("w:tblHeader")
                cells[0]._tc.getparent().get_or_add_trPr().append(repeat)


def convert(source: Path, destination: Path, cancel: Event) -> list[Notice]:
    try:
        text = source.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ConversionError("invalid_utf8", source.name) from exc
    renderer = Renderer(source, cancel)
    tree = SyntaxTreeNode(parser().parse(text))
    renderer.blocks(tree.children)
    renderer.checkpoint()
    renderer.document.save(destination)
    return list(dict.fromkeys(renderer.notices))
